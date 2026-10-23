# -*- coding: utf-8 -*-
"""
Created on Tue Oct 20 20:26:17 2026

@author: Carlos
"""

"""
Reaction Stoichiometry & Lab Notebook Module
Allows drawing full chemical reactions, robust PubChem IUPAC name fetching via InChIKey,
fully synchronized dynamic state, clean reaction schemes with condition overlay (above/below arrow),
Word (.docx) export, and dual file export.
"""

import streamlit as st
import pandas as pd
import urllib.request
import json
import io
from datetime import datetime

# --- Dependency Verification ---
try:
    from rdkit import Chem
    from rdkit.Chem import Descriptors
    from rdkit.Chem.rdMolDescriptors import CalcMolFormula
    RDKIT_AVAILABLE = True
except ImportError:
    RDKIT_AVAILABLE = False

try:
    from streamlit_ketcher import st_ketcher
    KETCHER_AVAILABLE = True
except ImportError:
    KETCHER_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Importar el motor de composición de Reaction_Scheme_Exporter
try:
    from Reaction_Scheme_Exporter import make_scheme, to_png_bytes
    SCHEME_EXPORTER_AVAILABLE = True
except ImportError:
    SCHEME_EXPORTER_AVAILABLE = False


def add_page_number_to_footer(run):
    """Añade el campo de número de página en un run del pie de página de docx."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def get_pubchem_iupac_by_inchikey(inchikey_str):
    """Consulta el nombre IUPAC oficial en PubChem usando el InChIKey."""
    if not inchikey_str:
        return None
    try:
        url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/inchikey/{inchikey_str}/property/IUPACName/json"
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=4) as response:
            data = json.loads(response.read().decode())
            props = data.get("PropertyTable", {}).get("Properties", [])
            if props:
                return props[0].get("IUPACName", None)
    except Exception:
        pass
    return None


def draw_reaction_scheme(rxn_smiles, catalyst="", solvent="", temperature=""):
    """Genera una imagen vectorial/PIL limpia de la reacción con las condiciones sobre/bajo la flecha."""
    if not SCHEME_EXPORTER_AVAILABLE:
        return None
    try:
        res = make_scheme(
            smiles_str=rxn_smiles,
            width_cm=16.0,
            dpi=600,
            fit_width=True,
            mode="reaction",
            solvent=solvent,
            catalyst=catalyst,
            temperature=temperature,
            add_labels=False
        )
        if res and res.get("image"):
            return to_png_bytes(res["image"], res["dpi"])
    except Exception:
        pass
    return None


def update_exp_key():
    """Callback para actualizar la clave del experimento dinámicamente cuando cambia el usuario."""
    user_name = st.session_state.get("dyn_user_name", "")
    if user_name.strip():
        words = user_name.strip().split()
        initials = "".join([w[0].upper() for w in words if w])
        if initials:
            st.session_state.dyn_exp_key = f"{initials}-001"
        else:
            st.session_state.dyn_exp_key = "EXP-001"
    else:
        st.session_state.dyn_exp_key = "EXP-001"


def generate_docx_document(experiment_key, user_name, solvent, catalyst, temperature, reaction_png_bytes, editable_components, solvent_row):
    """Genera el objeto BytesIO con el documento Word listo para descarga."""
    doc = Document()
    
    # Configurar Tamaño A4 y Márgenes (2,5 cm)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    printable_width = Cm(16.0)

    # Encabezado
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_text = f"Libreta de laboratorio de {user_name.strip()}" if user_name.strip() else "Libreta de laboratorio"
    run_hdr = p_hdr.add_run(header_text)
    run_hdr.font.size = Pt(9)
    run_hdr.font.italic = True
    run_hdr.font.color.rgb = RGBColor(128, 128, 128)

    # Pie de página
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_ftr = p_ftr.add_run()
    run_ftr.font.size = Pt(9)
    run_ftr.font.color.rgb = RGBColor(128, 128, 128)
    add_page_number_to_footer(run_ftr)

    # Título: Clave del experimento
    p_title = doc.add_paragraph(style='Heading 2')
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run(experiment_key)
    run_title.bold = True
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor(0, 51, 102)
    p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Fecha
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(6)
    current_date_str = datetime.now().strftime("%d/%m/%Y")
    run_date = p_date.add_run(current_date_str)
    run_date.font.size = Pt(11)
    run_date.font.color.rgb = RGBColor(100, 100, 100)
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Esquema de Reacción con las condiciones compuestas en la imagen
    if reaction_png_bytes:
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(8)
        
        image_stream = io.BytesIO(reaction_png_bytes)
        doc.add_picture(image_stream, width=printable_width)

    p_tab_title = doc.add_paragraph()
    p_tab_title.paragraph_format.space_before = Pt(6)
    p_tab_title.paragraph_format.space_after = Pt(4)
    p_tab_title.add_run("Tabla Estequiométrica de Preparación:").bold = True

    # Tabla
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    headers = ["Rol", "Fórmula", "Nombre", "PM (g/mol)", "Densidad (g/mL)", "Eq.", "mmol", "Masa/Vol"]
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9)

    # Filas
    for comp in editable_components:
        row_cells = table.add_row().cells
        vol_str = f"{comp['Volumen (mL)']} mL" if comp['Volumen (mL)'] > 0 else f"{int(comp['Masa (mg)'])} mg"
        row_values = [
            str(comp["Rol"]),
            str(comp["Fórmula"]),
            str(comp["Nombre (Editable)"]),
            str(comp["PM (g/mol)"]),
            str(comp["Densidad (g/mL)"]),
            str(comp["Equiv."]),
            str(comp["mmol"]),
            vol_str
        ]
        for i, val in enumerate(row_values):
            row_cells[i].text = val
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if p.runs:
                p.runs[0].font.size = Pt(9)

    # Disolvente
    row_cells = table.add_row().cells
    solv_values = [
        solvent_row["Rol"],
        solvent_row["Fórmula"],
        solvent_row["Nombre (Editable)"],
        solvent_row["PM (g/mol)"],
        solvent_row["Densidad (g/mL)"],
        solvent_row["Equiv."],
        solvent_row["mmol"],
        f"{solvent_row['Volumen (mL)']} mL"
    ]
    for i, val in enumerate(solv_values):
        row_cells[i].text = val
        p = row_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if p.runs:
            p.runs[0].font.size = Pt(9)

    p_proc_title = doc.add_paragraph()
    p_proc_title.paragraph_format.space_before = Pt(12)
    p_proc_title.paragraph_format.space_after = Pt(4)
    p_proc_title.add_run("Procedimiento y Observaciones:").bold = True
    
    p_tab_title = doc.add_paragraph()
    p_tab_title.paragraph_format.space_before = Pt(6)
    p_tab_title.paragraph_format.space_after = Pt(4)

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer.getvalue()


def render_reaction_lab_notebook():
    st.title("🧪 Libreta de Laboratorio: Estequiometría e Informes")
    st.markdown(
        "Dibuja la **reacción química completa** en el lienzo. El sistema extraerá automáticamente los compuestos, "
        "consultará sus nombres IUPAC oficiales en PubChem mediante InChIKey y generará tu hoja de trabajo exportable a Word."
    )

    if not RDKIT_AVAILABLE or not KETCHER_AVAILABLE:
        st.warning("Se requieren RDKit y `streamlit-ketcher` para ejecutar este módulo.")
        return

    st.markdown("### ✍️ Esquema de Reacción")
    
    reaction_smiles = ""
    try:
        reaction_smiles = st_ketcher("", height=520)
    except Exception:
        reaction_smiles = ""

    if not reaction_smiles or not reaction_smiles.strip():
        st.info("Dibuja la reacción en el lienzo superior y pulsa **Apply** para comenzar.")
        return

    # --- CONTROL DE VERSIÓN Y RECALCULO FORZADO DE LA TABLA ---
    if "last_reaction" not in st.session_state:
        st.session_state.last_reaction = reaction_smiles
        st.session_state.rxn_version = 0

    if st.session_state.last_reaction != reaction_smiles:
        st.session_state.last_reaction = reaction_smiles
        st.session_state.rxn_version += 1  # Incrementamos la versión para invalidar las claves de los widgets
        
        # Limpiar estados dinámicos anteriores
        keys_to_delete = [k for k in st.session_state.keys() if k.startswith("dyn_")]
        for k in keys_to_delete:
            del st.session_state[k]

    version_suffix = f"_v{st.session_state.rxn_version}"

    reactants_list = []
    products_list = []

    try:
        if ">>" in reaction_smiles:
            parts = reaction_smiles.split(">>")
            raw_reactants = parts[0].split(".")
            raw_products = parts[1].split(".") if len(parts) > 1 else []
            
            for r_smi in raw_reactants:
                if r_smi.strip():
                    mol = Chem.MolFromSmiles(r_smi.strip())
                    if mol:
                        reactants_list.append(mol)
            for p_smi in raw_products:
                if p_smi.strip():
                    mol = Chem.MolFromSmiles(p_smi.strip())
                    if mol:
                        products_list.append(mol)
        elif ">" in reaction_smiles:
            parts = reaction_smiles.split(">")
            raw_reactants = parts[0].split(".")
            raw_products = parts[-1].split(".") if len(parts) > 1 else []
            
            for r_smi in raw_reactants:
                if r_smi.strip():
                    mol = Chem.MolFromSmiles(r_smi.strip())
                    if mol:
                        reactants_list.append(mol)
            for p_smi in raw_products:
                if p_smi.strip():
                    mol = Chem.MolFromSmiles(p_smi.strip())
                    if mol:
                        products_list.append(mol)
        else:
            mol = Chem.MolFromSmiles(reaction_smiles)
            if mol:
                reactants_list.append(mol)
    except Exception as e:
        st.error(f"No se pudo interpretar la reacción dibujada: {e}")
        return

    if not reactants_list and not products_list:
        st.warning("No se han detectado moléculas válidas.")
        return

    if "dyn_exp_key" not in st.session_state:
        st.session_state.dyn_exp_key = "EXP-001"

    st.markdown("---")
    st.markdown("### 📋 Datos Generales del Experimento")

    with st.container(border=True):
        col_user, col_exp = st.columns(2)
        with col_user:
            user_name = st.text_input(
                "Nombre del usuario (opcional):", 
                value="", 
                key="dyn_user_name", 
                on_change=update_exp_key
            )

        with col_exp:
            experiment_key = st.text_input(
                "Clave del Experimento:", 
                key="dyn_exp_key"
            )

    st.markdown("### ⚙️ Condiciones de Reacción y Parámetros")

    with st.container(border=True):
        col_c1, col_c2, col_c3 = st.columns(3)
        with col_c1:
            catalyst = st.text_input("Catalizador / Reactivos (sobre la flecha):", value="N/A", key="dyn_catalyst")
        with col_c2:
            solvent = st.text_input("Disolvente principal (bajo la flecha):", value="Etanol", key="dyn_solvent")
        with col_c3:
            temperature = st.text_input("Temperatura / Tiempo (bajo la flecha):", value="Reflujo, 2 h", key="dyn_temp")

    # Generar imagen mejorada de la reacción
    reaction_png_bytes = draw_reaction_scheme(
        rxn_smiles=reaction_smiles,
        catalyst=catalyst,
        solvent=solvent,
        temperature=temperature
    )
    
    if reaction_png_bytes:
        st.markdown("#### 🖼️ Previsualización del Esquema de Reacción")
        caption_text = f"Experimento: {experiment_key}"
        if user_name.strip():
            caption_text = f"Usuario: {user_name} | " + caption_text
        st.image(reaction_png_bytes, caption=caption_text, use_container_width=True)

    # Procesar componentes consultando PubChem mediante InChIKey
    components_data = []

    for mol in reactants_list:
        formula = CalcMolFormula(mol)
        mw = Descriptors.MolWt(mol)
        smi = Chem.MolToSmiles(mol, canonical=True)
        inchikey = Chem.MolToInchiKey(mol)
        
        iupac_name = get_pubchem_iupac_by_inchikey(inchikey)
        default_name = iupac_name if iupac_name else f"Reactivo ({formula})"
        
        components_data.append({
            "smiles": smi,
            "role": "Reactivo",
            "formula": formula,
            "mw": mw,
            "name": default_name,
            "density": 1.0
        })

    for mol in products_list:
        formula = CalcMolFormula(mol)
        mw = Descriptors.MolWt(mol)
        smi = Chem.MolToSmiles(mol, canonical=True)
        inchikey = Chem.MolToInchiKey(mol)
        
        iupac_name = get_pubchem_iupac_by_inchikey(inchikey)
        default_name = iupac_name if iupac_name else f"Producto ({formula})"
        
        components_data.append({
            "smiles": smi,
            "role": "Producto",
            "formula": formula,
            "mw": mw,
            "name": default_name,
            "density": 1.0
        })

    if not reactants_list:
        st.warning("Se detectaron productos pero ningún reactivo válido.")
        return

    st.markdown("---")
    st.markdown("### 📊 Tabla Estequiometría de Laboratorio (Editable)")
    st.info("💡 Los nombres se obtienen de PubChem vía InChIKey. Marca la casilla 'Sólido' para reactivos en estado sólido (el volumen no se calculará). Los productos tampoco calculan volumen.")

    col_ref1, col_ref2 = st.columns(2)
    with col_ref1:
        ref_options = [f"{c['formula']} - {c['name'][:35]}" for c in components_data if c["role"] == "Reactivo"]
        selected_ref_label = st.selectbox("Reactivo de Referencia (Limitante):", ref_options, key=f"dyn_ref_select{version_suffix}")
        selected_ref_formula = selected_ref_label.split(" - ")[0] if selected_ref_label else ""
    with col_ref2:
        ref_mmol = st.number_input("Milimoles (mmol) de Referencia:", min_value=0.01, value=10.0, step=1.0, key=f"dyn_ref_mmol{version_suffix}")

    table_rows = []
    editable_components = []

    for idx, comp in enumerate(components_data):
        is_ref = comp["role"] == "Reactivo" and comp["formula"] == selected_ref_formula
        
        # Incorporamos el sufijo de versión en cada clave (key) para forzar la recreación del widget
        widget_key_suffix = f"{version_suffix}_{idx}"
        
        if comp["role"] == "Reactivo":
            col_name, col_solid, col_dens, col_eq = st.columns([2, 1, 1, 1])
            with col_name:
                custom_name = st.text_input(f"Nombre ({comp['formula']})", value=comp["name"], key=f"dyn_name{widget_key_suffix}")
            with col_solid:
                is_solid = st.checkbox("Sólido", value=False, key=f"dyn_solid{widget_key_suffix}")
            with col_dens:
                custom_dens = st.number_input("Dens. (g/mL)", value=float(comp["density"]), format="%.3f", step=0.001, key=f"dyn_dens{widget_key_suffix}")
            with col_eq:
                eq = 1.0 if is_ref else st.number_input(f"Eq. {comp['formula']}", value=1.0, step=0.1, key=f"dyn_eq{widget_key_suffix}")
                mmol = ref_mmol * eq
            
            mg = mmol * comp["mw"]
            ml = 0.0 if is_solid else ((mg / (custom_dens * 1000)) if custom_dens > 0 else 0.0)
        else:
            col_name, col_dens, col_eq = st.columns([2, 1, 1])
            with col_name:
                custom_name = st.text_input(f"Nombre ({comp['formula']})", value=comp["name"], key=f"dyn_name{widget_key_suffix}")
            with col_dens:
                custom_dens = st.number_input("Dens. (g/mL)", value=float(comp["density"]), format="%.3f", step=0.001, key=f"dyn_dens{widget_key_suffix}")
            with col_eq:
                eq = 1.0
                mmol = ref_mmol
            
            mg = mmol * comp["mw"]
            ml = 0.0

        row_dict = {
            "Rol": comp["role"],
            "Fórmula": comp["formula"],
            "Nombre (Editable)": custom_name,
            "PM (g/mol)": round(comp["mw"], 2),
            "Densidad (g/mL)": f"{custom_dens:.3f}",
            "Equiv.": eq,
            "mmol": round(mmol, 2),
            "Masa (mg)": round(mg),
            "Volumen (mL)": round(ml, 3)
        }
        table_rows.append(row_dict)
        editable_components.append(row_dict)

    # Disolvente
    st.markdown("#### 🧪 Disolvente de Reacción")
    solvent_label = solvent.strip() if solvent.strip() else "Disolvente"
    col_solv_vol, col_solv_dummy = st.columns([1, 2])
    with col_solv_vol:
        solvent_vol = st.number_input(f"Volumen de {solvent_label} (mL):", min_value=0.0, value=10.0, step=1.0, key=f"dyn_solvent_vol{version_suffix}")

    solvent_row = {
        "Rol": "Disolvente",
        "Fórmula": "-",
        "Nombre (Editable)": solvent_label,
        "PM (g/mol)": "-",
        "Densidad (g/mL)": "-",
        "Equiv.": "-",
        "mmol": "-",
        "Masa (mg)": "-",
        "Volumen (mL)": f"{solvent_vol:.1f}"
    }
    table_rows.append(solvent_row)

    df_lab = pd.DataFrame(table_rows)
    st.dataframe(df_lab, use_container_width=True)

    if not DOCX_AVAILABLE:
        st.warning("La librería `python-docx` no está instalada.")
        return

    st.markdown("---")
    st.markdown("### 📄 Exportar Informe a Word y Estructura")

    save_editable_format = st.checkbox(
        "Mostrar botón para descargar la estructura química editable (formato Ketcher/SMILES)", 
        value=True, 
        key=f"dyn_save_editable{version_suffix}"
    )

    docx_bytes = generate_docx_document(
        experiment_key=experiment_key,
        user_name=user_name,
        solvent=solvent,
        catalyst=catalyst,
        temperature=temperature,
        reaction_png_bytes=reaction_png_bytes,
        editable_components=editable_components,
        solvent_row=solvent_row
    )

    if save_editable_format:
        col_btn1, col_btn2 = st.columns(2)
        with col_btn1:
            st.download_button(
                label="📄 Descargar Documento Word (.docx)",
                data=docx_bytes,
                file_name=f"Libreta_Laboratorio_{experiment_key}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True
            )
        with col_btn2:
            st.download_button(
                label="🧪 Descargar Estructura Editable (.smi)",
                data=reaction_smiles.encode("utf-8"),
                file_name=f"Libreta_Laboratorio_{experiment_key}.smi",
                mime="text/plain",
                type="secondary",
                use_container_width=True,
                help="Guarda la cadena de reacción. Puedes volver a cargarla en Ketcher cuando lo necesites."
            )
    else:
        st.download_button(
            label="📄 Descargar Documento Word (.docx)",
            data=docx_bytes,
            file_name=f"Libreta_Laboratorio_{experiment_key}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            use_container_width=True
        )